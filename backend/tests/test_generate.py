"""课件渲染测试（纯本地，不依赖 LLM）。"""

from pathlib import Path

from docx import Document
from pptx import Presentation

from app.generate.creative import save_html
from app.generate.exam import render_exam
from app.generate.ppt import render_ppt
from app.generate.word import render_word


def test_render_ppt(tmp_path):
    slides = [
        {"title": "封面", "points": ["TCP 三次握手", "大二 · 45分钟"]},
        {"title": "内容页", "points": ["要点1", "要点2"]},
    ]
    path = render_ppt(slides, str(tmp_path / "t.pptx"))
    prs = Presentation(path)
    assert len(prs.slides) == 2
    title_shape = prs.slides[0].shapes.title
    assert title_shape is not None
    assert title_shape.text == "封面"


def test_render_exam(tmp_path):
    questions = [
        {"type": "选择", "content": "三次握手用于？", "answer": "建立连接"},
        {"type": "简答", "content": "简述三次握手过程。", "answer": "SYN → SYN+ACK → ACK"},
    ]
    path = render_exam(questions, str(tmp_path / "t.docx"))
    doc = Document(path)
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "试卷"
    assert "答案解析" in texts


def test_render_exam_sections_and_knowledge_annotations(tmp_path):
    """试卷两部分：试题 + 答案解析；试题行带考查知识点标注。"""
    questions = [
        {
            "type": "选择",
            "content": "三次握手用于？",
            "options": ["A. 2 次", "B. 3 次"],
            "answer": "B",
            "analysis": "握手用于同步初始序号",
            "knowledge_point": "TCP三次握手",
        },
        {
            "type": "填空",
            "content": "断开连接需____次挥手。",
            "answer": "四",
            "knowledge_point": "TCP四次挥手",
        },
    ]
    path = render_exam(questions, str(tmp_path / "t.docx"))
    texts = [p.text for p in Document(path).paragraphs]
    q_idx, a_idx = texts.index("试题"), texts.index("答案解析")
    assert q_idx < a_idx  # 试题在前，答案解析在后
    assert "TCP三次握手" in texts[q_idx + 1]  # 试题行带知识点标注
    assert "A. 2 次" in texts[q_idx + 2]  # 选择题选项单独成行
    assert "解析" in texts[a_idx + 1]  # 答案解析含解析说明


def test_render_word(tmp_path):
    data = {
        "objectives": {"knowledge": ["理解握手"], "ability": [], "emotion": []},
        "key_points": ["三次握手"],
        "difficult_points": [],
        "process": [{"stage": "导入", "minutes": 5, "content": "引入连接概念"}],
        "activities": ["课堂提问"],
        "homework": ["复习握手"],
    }
    path = render_word(data, str(tmp_path / "t.docx"))
    doc = Document(path)
    assert len(doc.paragraphs) > 5
    assert doc.paragraphs[0].text == "教案"


def test_render_ppt_default_names_unique():
    """连续两次默认落盘：文件名不同，互不覆盖。"""
    slides = [{"title": "封面", "points": ["要点"]}]
    p1 = render_ppt(slides)
    p2 = render_ppt(slides)
    assert p1 != p2
    assert Path(p1).is_file() and Path(p2).is_file()


def test_render_word_default_names_unique():
    """连续两次默认落盘：文件名不同，互不覆盖。"""
    data = {
        "key_points": ["三次握手"],
        "difficult_points": [],
        "process": [],
        "activities": [],
        "homework": [],
    }
    p1 = render_word(data)
    p2 = render_word(data)
    assert p1 != p2
    assert Path(p1).is_file() and Path(p2).is_file()


def test_render_exam_default_names_unique():
    """连续两次默认落盘：文件名不同，互不覆盖。"""
    questions = [{"type": "选择", "content": "q", "answer": "a"}]
    p1 = render_exam(questions)
    p2 = render_exam(questions)
    assert p1 != p2
    assert Path(p1).is_file() and Path(p2).is_file()


def test_save_html_default_names_unique():
    """连续两次默认落盘：文件名不同，互不覆盖。"""
    p1 = save_html("<html><body>1</body></html>")
    p2 = save_html("<html><body>2</body></html>")
    assert p1 != p2
    assert Path(p1).is_file() and Path(p2).is_file()
