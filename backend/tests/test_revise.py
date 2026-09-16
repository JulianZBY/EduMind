"""课件/教案迭代接口测试（stub 隔离）。"""

from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

import app.api.v1.revise as revise_module
from app.main import app

client = TestClient(app)


def test_revise(monkeypatch):
    async def fake_revise(slides, feedback):
        return [{"title": "修改后", "points": ["新要点"]}]

    monkeypatch.setattr(revise_module, "revise_ppt", fake_revise)
    r = client.post(
        "/api/v1/revise",
        json={"slides": [{"title": "t", "points": ["p"]}], "feedback": "简化"},
    )
    assert r.status_code == 200
    body = r.json()
    assert body["slides"][0]["title"] == "修改后"
    assert body["filename"].endswith(".pptx")


def test_revise_word(monkeypatch):
    """教案修改：与课件同模式——完整结构交 LLM 重排后重新渲染落盘。"""

    async def fake_revise(word, feedback):
        return {"key_points": ["修改后重点"], "difficult_points": [], "process": []}

    monkeypatch.setattr(revise_module, "revise_word", fake_revise)
    r = client.post(
        "/api/v1/revise/word",
        json={"word": {"key_points": ["原重点"]}, "feedback": "重点改为动手实践"},
    )
    assert r.status_code == 200
    body = r.json()
    assert body["word"]["key_points"] == ["修改后重点"]
    assert body["filename"].endswith(".docx")


def test_revise_word_consecutive_progressive(monkeypatch):
    """连续两次修改：各自返回新文件且内容递进，互不覆盖。"""

    rounds = [
        {"key_points": ["第一轮重点"], "homework": ["练习一"]},
        {"key_points": ["第二轮重点"], "homework": ["练习一", "拓展思考"]},
    ]

    async def fake_revise(word, feedback):
        return rounds[len(rounds) - 1] if "再改" in feedback else rounds[0]

    monkeypatch.setattr(revise_module, "revise_word", fake_revise)
    r1 = client.post(
        "/api/v1/revise/word", json={"word": {"key_points": ["原重点"]}, "feedback": "改重点"}
    )
    r2 = client.post(
        "/api/v1/revise/word",
        json={"word": r1.json()["word"], "feedback": "再改一次"},
    )
    assert r1.status_code == 200 and r2.status_code == 200
    f1, f2 = r1.json()["filename"], r2.json()["filename"]
    assert f1 != f2
    p1, p2 = Path("data/output") / f1, Path("data/output") / f2
    assert p1.is_file() and p2.is_file()
    texts1 = [p.text for p in Document(str(p1)).paragraphs]
    texts2 = [p.text for p in Document(str(p2)).paragraphs]
    # 内容递进：第二轮文件含新增作业；第一轮文件保持第一轮内容，未被覆盖
    assert not any("拓展思考" in t for t in texts1)
    assert any("拓展思考" in t for t in texts2)
    assert any("练习一" in t for t in texts1)
