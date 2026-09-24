"""参考资料关联与溯源测试（HTTP API 主接缝：stub 网关 + 直接种子数据）。

覆盖 ticket #4 验收项：
- 上传支持标记参考资料；备课请求携带参考资料标识
- 加权生效：相同查询下参考文档片段排名靠前（种子数据断言）
- 备课回复文案列出命中的来源文档名
- 教案产物回读含「参考资料」节及来源名
- 不带参考资料标识时行为与现状兼容
"""

import uuid

from docx import Document as DocxDocument
from fastapi.testclient import TestClient

import app.api.v1.chat as chat_module
import app.api.v1.documents as documents_module
import app.core.embedding.factory as embedding_factory_module
import app.core.orchestrator as orchestrator_module
import app.generate.outline as outline_module
import app.generate.ppt as ppt_module
import app.generate.word as word_module
import app.knowledge.vector_store as vector_store_module
from app.core.embedding.stub import StubEmbedder
from app.core.intent import TeachingIntent
from app.core.llm.base import ChatResult, LLMProvider
from app.db import SessionLocal, init_db
from app.db.models import Document
from app.knowledge.vector_store import VectorStore
from app.main import app

client = TestClient(app)
init_db()  # 幂等：确保表、列与默认用户存在

# 三个生成器（PPT/Word/提纲）共用的假响应：结构可解析，渲染走默认值
_FAKE_CHAT_JSON = (
    '{"slides": [{"title": "s", "points": ["p"]}], "key_points": ["k"], '
    '"difficult_points": [], "process": [], "activities": [], "homework": []}'
)


class RecordingProvider(LLMProvider):
    """捕获全部 chat prompt 的假网关：可断言传给 LLM 的上下文内容与排名。"""

    def __init__(self) -> None:
        self.prompts: list[str] = []

    async def chat(self, messages, **kwargs) -> ChatResult:
        self.prompts.append(messages[-1].content)
        return ChatResult(content=_FAKE_CHAT_JSON)

    async def embed(self, texts: list[str]) -> list[list[float]]:
        # 与 stub 同构：以文本长度为特征，query「导数」→ [2.0]*8
        return [[float(len(t))] * 8 for t in texts]


def _install(monkeypatch, provider: RecordingProvider, store: VectorStore | None) -> None:
    """伪装意图分析 + 全部 get_llm 引用 + 向量库（orchestrator 延迟 import，逐模块替换）。"""

    async def fake_analyze(text):
        return TeachingIntent(
            topic="导数", duration_minutes=45, style="学术", objectives=["a"], key_points=["b"]
        )

    # orchestrate 内部引用自己模块的 analyze_intent（早绑定），必须一并替换
    monkeypatch.setattr(chat_module, "analyze_intent", fake_analyze)
    monkeypatch.setattr(orchestrator_module, "analyze_intent", fake_analyze)
    monkeypatch.setattr(embedding_factory_module, "get_embedder", lambda: StubEmbedder())
    monkeypatch.setattr(ppt_module, "get_llm", lambda: provider)
    monkeypatch.setattr(word_module, "get_llm", lambda: provider)
    monkeypatch.setattr(outline_module, "get_llm", lambda: provider)
    monkeypatch.setattr(vector_store_module, "VectorStore", lambda: store)


def _seed_document(filename: str, is_reference: bool) -> str:
    """直接种子 Document 表，返回文档 id（供向量库 doc_id 与文档名映射对齐）。"""
    db = SessionLocal()
    try:
        doc = Document(
            user_id="default",
            filename=filename,
            file_path=f"data/uploads/{filename}",
            file_type="pdf",
            status="已完成",
            is_reference=is_reference,
        )
        db.add(doc)
        db.commit()
        return doc.id
    finally:
        db.close()


def test_upload_mark_and_default_reference(monkeypatch):
    """上传可标记参考资料；不标记时默认 False；列表接口回显标记。"""

    async def fake_parse(doc_id: str) -> str:
        return "parsed"

    monkeypatch.setattr(documents_module, "parse_document", fake_parse)

    files = {"file": (f"参考讲义_{uuid.uuid4().hex[:6]}.pdf", b"%PDF-1.4 fake", "application/pdf")}
    marked = client.post("/api/v1/documents/upload", files=files, data={"is_reference": "true"})
    assert marked.status_code == 200
    assert marked.json()["is_reference"] is True

    plain = client.post(
        "/api/v1/documents/upload",
        files={
            "file": (f"普通资料_{uuid.uuid4().hex[:6]}.pdf", b"%PDF-1.4 fake", "application/pdf")
        },
    )
    assert plain.status_code == 200
    assert plain.json()["is_reference"] is False

    listed = {d["id"]: d for d in client.get("/api/v1/documents").json()["documents"]}
    assert listed[marked.json()["id"]]["is_reference"] is True
    assert listed[plain.json()["id"]]["is_reference"] is False


def test_search_ranks_reference_chunks_first(tmp_path):
    """种子数据断言：相同查询下，参考文档片段加权后排名靠前；未加权时靠后。"""
    store = VectorStore(str(tmp_path / "v.db"))
    # 普通片段天然更近（0.9√8 < 1.0√8）；加权折扣 0.5 后参考片段（0.5√8）反超
    store.add("normal_doc", ["普通知识点片段"], [[0.9] * 8])
    store.add("ref_doc", ["参考讲义片段"], [[1.0] * 8])
    query = [0.0] * 8

    plain = store.search(query, k=2)
    assert [h["doc_id"] for h in plain] == ["normal_doc", "ref_doc"]

    boosted = store.search(query, k=2, boost_doc_ids={"ref_doc"})
    assert [h["doc_id"] for h in boosted] == ["ref_doc", "normal_doc"]
    assert boosted[0]["distance"] > boosted[1]["distance"]  # 返回原始距离，仅排名生效


def test_prep_ranks_reference_chunk_first(monkeypatch, tmp_path):
    """备课请求携带参考资料标识后，参考文档片段在生成上下文中排到普通片段之前。"""
    suffix = uuid.uuid4().hex[:6]
    store = VectorStore(str(tmp_path / "v.db"))
    # 普通片段天然更近（0.9√8 < 1.0√8），参考片段加权后（0.5√8）反超居首
    store.add(f"normal_{suffix}", [f"普通片段内容_{suffix}"], [[2.9] * 8])
    store.add(f"ref_{suffix}", [f"参考片段内容_{suffix}"], [[3.0] * 8])
    provider = RecordingProvider()
    _install(monkeypatch, provider, store)

    r = client.post(
        "/api/v1/chat",
        json={
            "messages": [{"role": "user", "content": "讲导数"}],
            "reference_doc_ids": [f"ref_{suffix}"],
        },
    )

    assert r.status_code == 200
    assert r.json()["artifacts"] is not None
    ppt_prompt = next(p for p in provider.prompts if "课件设计师" in p)
    assert ppt_prompt.index(f"参考片段内容_{suffix}") < ppt_prompt.index(f"普通片段内容_{suffix}")


def test_reply_and_artifacts_list_hit_source_doc_names(monkeypatch, tmp_path):
    """携带参考资料标识时，回复文案与 artifacts 列出本次命中的来源文档名。"""
    suffix = uuid.uuid4().hex[:6]
    ref_name = f"参考讲义_{suffix}.pdf"
    normal_name = f"背景资料_{suffix}.pdf"
    ref_id = _seed_document(ref_name, is_reference=True)
    normal_id = _seed_document(normal_name, is_reference=False)

    store = VectorStore(str(tmp_path / "v.db"))
    store.add(ref_id, [f"参考讲义片段_{suffix}"], [[3.0] * 8])
    store.add(normal_id, [f"背景资料片段_{suffix}"], [[2.9] * 8])
    provider = RecordingProvider()
    _install(monkeypatch, provider, store)

    r = client.post(
        "/api/v1/chat",
        json={
            "messages": [{"role": "user", "content": "讲导数"}],
            "reference_doc_ids": [ref_id],
        },
    )

    assert r.status_code == 200
    body = r.json()
    assert ref_name in body["content"]
    assert normal_name in body["content"]
    assert set(body["artifacts"]["references"]) == {ref_name, normal_name}


def test_word_artifact_contains_reference_section(monkeypatch, tmp_path):
    """教案产物回读：自动追加「参考资料」一节，且节内列出来源文档名。"""
    suffix = uuid.uuid4().hex[:6]
    ref_name = f"参考讲义_{suffix}.pdf"
    ref_id = _seed_document(ref_name, is_reference=True)

    store = VectorStore(str(tmp_path / "v.db"))
    store.add(ref_id, [f"参考讲义片段_{suffix}"], [[3.0] * 8])
    provider = RecordingProvider()
    _install(monkeypatch, provider, store)

    r = client.post(
        "/api/v1/chat",
        json={
            "messages": [{"role": "user", "content": "讲导数"}],
            "reference_doc_ids": [ref_id],
        },
    )

    assert r.status_code == 200
    word_path = r.json()["artifacts"]["word"]["path"]
    doc = DocxDocument(word_path)
    texts = [p.text for p in doc.paragraphs]
    heading_idx = next(i for i, t in enumerate(texts) if "参考资料" in t)  # 节标题存在
    assert any(ref_name in t for t in texts[heading_idx:])  # 节内含来源文档名


def test_chat_without_reference_ids_keeps_current_behavior(monkeypatch, tmp_path):
    """不带参考资料标识：检索照常生效，回复无溯源文案，教案无参考资料节。"""
    suffix = uuid.uuid4().hex[:6]
    store = VectorStore(str(tmp_path / "v.db"))
    store.add(f"hit_{suffix}", [f"普通命中片段_{suffix}"], [[2.9] * 8])
    provider = RecordingProvider()
    _install(monkeypatch, provider, store)

    r = client.post("/api/v1/chat", json={"messages": [{"role": "user", "content": "讲导数"}]})

    assert r.status_code == 200
    body = r.json()
    assert body["artifacts"] is not None
    assert "已融合本地知识库" in body["content"]  # 现有文案保持
    assert "来源文档" not in body["content"]  # 无溯源文案
    assert body["artifacts"]["references"] == []
    assert f"普通命中片段_{suffix}" in "\n".join(provider.prompts)  # 检索仍生效
    word_path = body["artifacts"]["word"]["path"]
    texts = [p.text for p in DocxDocument(word_path).paragraphs]
    assert not any("参考资料" in t for t in texts)  # 教案无参考资料节
