"""解析器单元测试（本地格式，不依赖云端）。"""

import asyncio

from docx import Document as DocxDocument
from pptx import Presentation

from app.knowledge.parsers.ppt import PptParser
from app.knowledge.parsers.word import WordParser


def test_word_parser(tmp_path):
    doc = DocxDocument()
    doc.add_paragraph("教学目标：理解网络分层")
    p = tmp_path / "test.docx"
    doc.save(str(p))
    result = asyncio.run(WordParser().parse(str(p)))
    assert "教学目标" in result


def test_ppt_parser(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "计算机网络概述"
    p = tmp_path / "test.pptx"
    prs.save(str(p))
    result = asyncio.run(PptParser().parse(str(p)))
    assert "计算机网络概述" in result


def test_get_parser_unknown():
    from app.knowledge.parsers import get_parser

    try:
        get_parser("xyz")
        assert False, "应抛出 ValueError"
    except ValueError:
        pass
