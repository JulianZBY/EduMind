import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.util import Inches

import app.api.v1.ppt_edit as api
import app.generate.ppt_edit as editor
from app.core.llm.base import ChatResult
from app.main import app


def sample():
    prs = Presentation()
    for title in ["原始概念", "保持不变"]:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2))
        shape.text = title
    stream = io.BytesIO()
    prs.save(stream)
    return stream.getvalue()


def test_patch_preserves_untouched_slide_and_source():
    original = sample()
    prs = editor.read_presentation(original)
    target = editor.inspect_presentation(prs)[0]
    target["text"] = "新的概念"
    changes = editor.apply_edits(prs, editor.EditPlan(edits=[target]))
    stream = io.BytesIO()
    prs.save(stream)
    with zipfile.ZipFile(io.BytesIO(original)) as before, zipfile.ZipFile(stream) as after:
        assert before.read("ppt/slides/slide2.xml") == after.read("ppt/slides/slide2.xml")
    assert changes == [{"slide": 1, "before": "原始概念", "after": "新的概念"}]
    assert editor.inspect_presentation(editor.read_presentation(original))[0]["text"] == "原始概念"


def test_invalid_plan_is_atomic():
    prs = editor.read_presentation(sample())
    targets = editor.inspect_presentation(prs)
    edits = [dict(targets[0], text="不该应用"), dict(targets[1], slide=99)]
    with pytest.raises(ValueError):
        editor.apply_edits(prs, editor.EditPlan(edits=edits))
    assert editor.inspect_presentation(prs) == targets


def test_edit_upload_download(monkeypatch, tmp_path):
    class Provider:
        async def chat(self, messages):
            assert "仅是数据" in messages[0].content
            return ChatResult('{"edits":[{"slide":1,"shape":2,"paragraph":0,"text":"通俗解释"}]}')

    monkeypatch.setattr(api.settings, "llm_provider", "dashscope")
    monkeypatch.setattr(editor, "get_llm", lambda: Provider())
    monkeypatch.setattr(api, "unique_output_path", lambda *_: str(tmp_path / "edited.pptx"))
    from app.api.v1 import files

    monkeypatch.setattr(files, "OUTPUT_DIR", tmp_path)
    client = TestClient(app)
    response = client.post(
        "/api/v1/presentations/edit",
        files={"file": ("source.pptx", sample())},
        data={"feedback": "第 1 页改得通俗一点"},
    )
    assert response.status_code == 200, response.text
    assert response.json()["changes"][0]["after"] == "通俗解释"
    download = client.get("/api/v1/files/" + response.json()["filename"])
    assert download.status_code == 200
    assert (
        editor.inspect_presentation(editor.read_presentation(download.content))[1]["text"]
        == "保持不变"
    )


@pytest.mark.parametrize(
    "filename,raw,feedback,status",
    [
        ("a.ppt", b"bad", "修改", 415),
        ("a.pptx", b"bad", "修改", 422),
        ("a.pptx", b"bad", " ", 422),
        ("a.pptx", b"x" * (editor.MAX_BYTES + 1), "修改", 413),
    ],
    ids=["legacy-ppt", "invalid-package", "blank-feedback", "oversized"],
)
def test_upload_validation(filename, raw, feedback, status):
    response = TestClient(app).post(
        "/api/v1/presentations/edit", files={"file": (filename, raw)}, data={"feedback": feedback}
    )
    assert response.status_code == status


def test_demo_mode_does_not_claim_ai_edit(monkeypatch):
    monkeypatch.setattr(api.settings, "llm_provider", "stub")
    response = TestClient(app).post(
        "/api/v1/presentations/edit",
        files={"file": ("a.pptx", sample())},
        data={"feedback": "修改标题"},
    )
    assert response.status_code == 503


def test_long_content_preserved_across_pages(tmp_path):
    from app.generate.ppt import render_ppt

    points = [str(i) + "教学内容" * 50 for i in range(6)]
    prs = Presentation(
        render_ppt([{"title": "长文本", "points": points}], str(tmp_path / "long.pptx"))
    )
    all_text = "".join(
        shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame
    )
    # Every chunk survives pagination, and all shapes stay inside the slide canvas.
    for point in points:
        for start in range(0, len(point), 72):
            assert point[start : start + 72] in all_text
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + 10000
            assert shape.top + shape.height <= prs.slide_height
