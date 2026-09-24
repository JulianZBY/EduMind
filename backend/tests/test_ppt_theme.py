"""课件主题化渲染测试（ticket #9）。

主接缝：HTTP API + stub 网关——生成结构含页面角色字段，排版规则（要点数/字数上限）不变。
副接缝：渲染器纯函数——pptx 回读断言角色版式差异、配色主题映射、旧结构回退默认版式。
"""

import uuid

from fastapi.testclient import TestClient
from pptx import Presentation
from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Pt

import app.core.conversation as conversation_module
import app.core.embedding.factory as embedding_factory_module
import app.generate.outline as outline_module
import app.generate.ppt as ppt_module
import app.generate.word as word_module
import app.knowledge.vector_store as vector_store_module
from app.core.embedding.stub import StubEmbedder
from app.core.intent import TeachingIntent
from app.core.llm.providers.stub import StubProvider
from app.generate.ppt import render_ppt
from app.knowledge.vector_store import VectorStore
from app.main import app

client = TestClient(app)

_ROLES = {"封面", "目录", "内容", "总结"}


class RecordingStub(StubProvider):
    """记录 prompt 的 stub 网关：断言 prompt 约定的同时保持 stub 响应可解析。"""

    def __init__(self) -> None:
        self.prompts: list[str] = []

    async def chat(self, messages, **kwargs):
        self.prompts.append(messages[-1].content if messages else "")
        return await super().chat(messages, **kwargs)


def _accent_rgb(slide) -> str | None:
    """带回读：封面色块/总结强调条的强调色（hex），无实心 RGB 填充图形则 None。"""
    for shape in slide.shapes:
        if shape.shape_type != MSO_SHAPE_TYPE.AUTO_SHAPE:
            continue
        fill = getattr(shape, "fill", None)
        if fill is None or fill.type != MSO_FILL.SOLID:
            continue
        if fill.fore_color.type != MSO_COLOR_TYPE.RGB:
            continue
        return str(fill.fore_color.rgb)
    return None


def _shape_with_text(slide, text):
    """带回读：找文本含 text 的形状（首个匹配）。"""
    for shape in slide.shapes:
        if shape.has_text_frame and text in shape.text_frame.text:
            return shape
    return None


def _max_font_size(shape) -> int:
    """带回读：形状内最大字号（EMU），无显式字号返回 0。"""
    sizes = [
        run.font.size
        for para in shape.text_frame.paragraphs
        for run in para.runs
        if run.font.size is not None
    ]
    return max(sizes, default=0)


# ---- 主接缝：prompt 约定 + stub 网关下的生成结构 ----


async def test_prompt_declares_role_and_stub_structure_satisfies(monkeypatch):
    """生成 prompt 约定页面角色字段；stub 网关响应满足约定（结构含合法角色字段）。"""
    provider = RecordingStub()
    monkeypatch.setattr(ppt_module, "get_llm", lambda: provider)

    slides = await ppt_module.generate_ppt_structure({"topic": "TCP"}, "TCP 知识")

    prompt = provider.prompts[0]
    assert "role" in prompt
    for role in _ROLES:
        assert role in prompt  # 角色取值进入 prompt 约定
    assert slides, "stub 下应返回可解析的 slides"
    assert {s["role"] for s in slides} == _ROLES  # 封面/目录/内容/总结齐备且取值合法


def test_chat_artifacts_slides_carry_roles_and_keep_typography(monkeypatch, tmp_path):
    """stub 网关走完整备课：产物 slides 每页带合法角色；要点数与字数上限不变。

    get_llm 逐模块替换（早绑定引用）+ 检索向量化走 Embedder 工厂引用：.env 即便配置真实
    provider 也确定性走 stub（spec：LLM 网关一律 stub 保证确定性）。
    """

    async def fake_analyze(text):
        return TeachingIntent(
            topic="TCP", duration_minutes=45, style="学术", objectives=["a"], key_points=["b"]
        )

    stub = StubProvider()
    # 意图分析住 core 状态机（票 05 收编）：伪装意图分析；orchestrator 只消费累积意图
    monkeypatch.setattr(conversation_module, "analyze_intent", fake_analyze)
    monkeypatch.setattr(embedding_factory_module, "get_embedder", lambda: StubEmbedder())
    monkeypatch.setattr(ppt_module, "get_llm", lambda: stub)
    monkeypatch.setattr(word_module, "get_llm", lambda: stub)
    monkeypatch.setattr(outline_module, "get_llm", lambda: stub)
    monkeypatch.setattr(vector_store_module, "VectorStore", lambda: _new_store(tmp_path))

    r = client.post("/api/v1/chat", json={"messages": [{"role": "user", "content": "讲TCP"}]})

    assert r.status_code == 200
    body = r.json()
    slides = body["artifacts"]["ppt"]["slides"]
    assert slides
    assert all(s.get("role") in _ROLES for s in slides)
    for s in slides:  # 排版规则保持：要点数与字数上限不变
        assert len(s["title"]) <= 15
        assert len(s["points"]) <= 5
        assert all(len(p) <= 30 for p in s["points"])
    # 风格偏好贯穿到渲染产物：封面色块为学术主题配色
    accent = _accent_rgb(Presentation(body["artifacts"]["ppt"]["path"]).slides[0])
    assert accent == ppt_module.PPT_THEMES["学术"]["accent"]


# ---- 副接缝：渲染器纯函数回读断言 ----


def test_render_role_layouts_distinguishable(tmp_path):
    """封面/目录/内容/总结按角色差异化版式，回读可区分。"""
    slides = [
        {"role": "封面", "title": "TCP 协议", "points": ["计算机网络", "大二课程"]},
        {"role": "目录", "title": "目录", "points": ["三次握手", "四次挥手", "滑动窗口"]},
        {"role": "内容", "title": "三次握手", "points": ["SYN", "SYN+ACK", "ACK"]},
        {"role": "总结", "title": "小结", "points": ["可靠连接", "流量控制"]},
    ]
    path = render_ppt(slides, str(tmp_path / "t.pptx"), style="学术")
    cover, toc, content, summary = Presentation(path).slides

    # 封面：无标题占位符 + 色块 + 大标题（≥40pt）
    assert cover.shapes.title is None
    assert _accent_rgb(cover) is not None
    cover_title = _shape_with_text(cover, "TCP 协议")
    assert cover_title is not None
    assert _max_font_size(cover_title) >= Pt(40)

    # 目录：分栏——首条与末条落在不同文本框
    first = _shape_with_text(toc, "三次握手")
    last = _shape_with_text(toc, "滑动窗口")
    assert first is not None and last is not None
    assert first is not last

    # 内容页：默认版式（标题 + 正文占位符）
    content_title = content.shapes.title
    assert content_title is not None
    assert content_title.text == "三次握手"

    # 总结：标题占位符强调色加粗 + 强调色条
    summary_title = summary.shapes.title
    assert summary_title is not None
    runs = summary_title.text_frame.paragraphs[0].runs
    assert runs and runs[0].font.bold
    assert str(runs[0].font.color.rgb) == ppt_module.PPT_THEMES["学术"]["accent"]
    assert _accent_rgb(summary) is not None


def test_theme_maps_style_preference(tmp_path):
    """三套配色映射风格偏好：简约=黑白灰、学术=深蓝、活泼=暖橙；未表达回退简约。"""

    def rgb(hex_str: str) -> tuple[int, int, int]:
        return int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16)

    accents: dict[str, str] = {}
    for i, style in enumerate(("简约", "学术", "活泼", "", "未知风格")):
        slides = [{"role": "封面", "title": "T", "points": ["p"]}]
        path = render_ppt(slides, str(tmp_path / f"t{i}.pptx"), style=style)
        accent = _accent_rgb(Presentation(path).slides[0])
        assert accent is not None
        accents[style] = accent

    gray, blue, warm = rgb(accents["简约"]), rgb(accents["学术"]), rgb(accents["活泼"])
    assert gray[0] == gray[1] == gray[2]  # 黑白灰：RGB 等值
    assert blue[2] > blue[0]  # 深蓝：蓝分量主导
    assert warm[0] > warm[2] and warm[0] > warm[1]  # 暖橙：红分量主导
    assert accents[""] == accents["简约"]  # 未表达风格 → 默认主题
    assert accents["未知风格"] == accents["简约"]  # 未知风格 → 默认主题


def test_render_without_role_falls_back_to_default_layout(tmp_path):
    """旧结构（缺失角色字段）不崩溃：回退默认版式（标题+正文占位符），无封面色块。"""
    slides = [
        {"title": "封面", "points": ["TCP 三次握手"]},
        {"title": "内容页", "points": ["要点1", "要点2"]},
        {"role": "未知", "title": "非法取值", "points": ["x"]},
        {"role": "封面页", "title": "带页后缀", "points": []},
    ]
    path = render_ppt(slides, str(tmp_path / "t.pptx"))
    s0, s1, s2, s3 = Presentation(path).slides

    for s in (s0, s1, s2):  # 缺失/非法角色 → 默认版式
        assert s.shapes.title is not None
        assert _accent_rgb(s) is None
    t0 = s0.shapes.title
    assert t0 is not None
    assert t0.text == "封面"
    assert s3.shapes.title is None  # 「封面页」容错归一化为封面版式
    assert _accent_rgb(s3) is not None


def _new_store(tmp_path):
    # 模块顶部已绑定真实 VectorStore 类：不受 monkeypatch 替换模块属性影响
    return VectorStore(str(tmp_path / f"vectors_{uuid.uuid4().hex[:6]}.db"))
