"""Word 解析：python-docx。"""

from docx import Document as DocxDocument

from app.knowledge.parsers.base import Parser


class WordParser(Parser):
    async def parse(self, file_path: str) -> str:
        doc = DocxDocument(file_path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
