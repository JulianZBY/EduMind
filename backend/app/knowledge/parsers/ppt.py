"""PPT 解析：python-pptx，抽取每页文本。"""

from pptx import Presentation

from app.knowledge.parsers.base import Parser


class PptParser(Parser):
    async def parse(self, file_path: str) -> str:
        prs = Presentation(file_path)
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
            slides.append(f"## 第{i}页\n" + "\n".join(texts))
        return "\n\n".join(slides)
