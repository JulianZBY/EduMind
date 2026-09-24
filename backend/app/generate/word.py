"""Word 教案生成：LLM 生成结构 + python-docx 渲染。"""

import json

from docx import Document as DocxDocument
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from app.core.llm.base import ChatMessage
from app.core.llm.factory import get_llm
from app.core.llm.parsing import parse_json
from app.generate import unique_output_path
from app.generate.document_style import style_document

_WORD_PROMPT = """你是教学设计专家。根据教学意图和知识内容，生成详细教案。

结构：
- objectives：教学目标（knowledge/ability/emotion 三组字符串数组）
- key_points：教学重点（字符串数组）
- difficult_points：教学难点（字符串数组）
- process：教学过程（数组，每项 {stage, minutes, content}）
- activities：课堂活动设计（字符串数组）
- homework：课后作业（字符串数组）

只输出 JSON，不要其他文字。

教学意图：__INTENT__
知识内容：__KNOWLEDGE__
"""


async def generate_word_structure(intent: dict, knowledge: str) -> dict:
    llm = get_llm()
    prompt = _WORD_PROMPT.replace("__INTENT__", json.dumps(intent, ensure_ascii=False)).replace(
        "__KNOWLEDGE__", knowledge[:6000]
    )
    result = await llm.chat([ChatMessage(role="user", content=prompt)])
    return parse_json(result.content)


def render_word(
    data: dict, output_path: str | None = None, references: list[str] | None = None
) -> str:
    """渲染 .docx，返回文件路径。references 非空时自动追加「参考资料」一节（溯源）。"""
    if output_path is None:
        output_path = unique_output_path("lesson_plan", ".docx")

    doc = DocxDocument()
    style_document(doc)
    label = doc.add_paragraph("EDUMIND  /  TEACHING NOTES")
    label.runs[0].font.size = Pt(9)
    label.runs[0].font.color.rgb = RGBColor.from_string("728577")
    doc.add_heading(data.get("title") or "课堂教学设计", level=0)
    doc.add_paragraph("教案 · 教学目标 / 课堂路径 / 学习活动", style="Subtitle")

    objectives = data.get("objectives", {})
    doc.add_heading("一、教学目标", level=1)
    if isinstance(objectives, list):
        doc.add_paragraph("教学目标：" + "；".join(objectives))
    else:
        for label, key in [
            ("知识目标", "knowledge"),
            ("能力目标", "ability"),
            ("情感目标", "emotion"),
        ]:
            items = objectives.get(key, [])
            if items:
                doc.add_paragraph(f"{label}：" + "；".join(items))

    doc.add_heading("二、教学重点与难点", level=1)
    doc.add_paragraph("重点：" + "；".join(data.get("key_points", [])))
    doc.add_paragraph("难点：" + "；".join(data.get("difficult_points", [])))

    doc.add_heading("三、教学过程", level=1)
    for i, step in enumerate(data.get("process", []), 1):
        heading = doc.add_paragraph(style="Heading 2")
        heading.add_run(f"{i:02}  {step.get('stage', '')}    /    {step.get('minutes', '?')} 分钟")
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "EFF4EC")
        heading._p.get_or_add_pPr().append(shading)
        doc.add_paragraph(str(step.get("content", "")))

    doc.add_heading("四、课堂活动设计", level=1)
    for act in data.get("activities", []):
        doc.add_paragraph(act, style="List Bullet")

    doc.add_heading("五、课后作业", level=1)
    for hw in data.get("homework", []):
        doc.add_paragraph(hw, style="List Bullet")

    if references:
        doc.add_heading("六、参考资料", level=1)
        for name in references:
            doc.add_paragraph(name, style="List Bullet")

    doc.save(output_path)
    return output_path
