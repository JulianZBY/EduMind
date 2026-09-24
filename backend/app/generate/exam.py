"""试卷生成：LLM 自编出题 + 题库入库 + python-docx 渲染。

按需生成入口（产物区一键生成）；组题模式与网络搜题仅预留（DESIGN.md §5.4）。
"""

import json
from typing import TYPE_CHECKING

from docx import Document as DocxDocument
from sqlalchemy import select

from app.core.llm.base import ChatMessage
from app.core.llm.factory import get_llm
from app.core.llm.parsing import parse_json
from app.generate import unique_output_path
from app.generate.document_style import style_document

if TYPE_CHECKING:
    from collections.abc import Sequence

    from app.db.models import KnowledgeNode

_EXAM_PROMPT = """你是出题专家。根据教学意图和知识内容，生成试卷。

要求：
- 生成 __N__ 道题，题型覆盖选择/填空/简答
- 每题含 type、content、answer、analysis（答案解析）、knowledge_point（本题考查的知识点，取自知识内容中的知识点名称）
- 选择题附 options（选项数组）

只输出 JSON，不要其他文字：
{"questions": [{"type": "选择", "content": "题目", "options": ["A. …", "B. …"], "answer": "答案", "analysis": "解析", "knowledge_point": "知识点"}]}

教学意图：__INTENT__
知识内容：__KNOWLEDGE__
"""


async def generate_exam(intent: dict, knowledge: str, n: int = 5) -> list[dict]:
    llm = get_llm()
    prompt = (
        _EXAM_PROMPT.replace("__N__", str(n))
        .replace("__INTENT__", json.dumps(intent, ensure_ascii=False))
        .replace("__KNOWLEDGE__", knowledge[:6000])
    )
    result = await llm.chat([ChatMessage(role="user", content=prompt)])
    data = parse_json(result.content)
    return data.get("questions", [])


def save_questions_to_bank(questions: list[dict], user_id: str = "default") -> int:
    """生成题目写入题库：来源=自编，并按考查知识点关联图谱节点（主考/涉及）。"""
    from app.db import SessionLocal
    from app.db.models import KnowledgeNode, Question, QuestionKnowledge

    if not questions:
        return 0
    db = SessionLocal()
    try:
        nodes = (
            db.execute(select(KnowledgeNode).where(KnowledgeNode.user_id == user_id))
            .scalars()
            .all()
        )
        saved = 0
        for q in questions:
            row = Question(
                user_id=user_id,
                content=str(q.get("content", "")).strip(),
                answer=str(q.get("answer", "")).strip(),
                type=str(q.get("type", "")).strip()[:20] or "简答",
                source_type="自编",
            )
            db.add(row)
            db.flush()  # 先取 row.id 再建题目-知识点关联
            for i, node in enumerate(
                _match_nodes(nodes, str(q.get("knowledge_point", "")).strip())
            ):
                db.add(
                    QuestionKnowledge(
                        question_id=row.id,
                        knowledge_id=node.id,
                        weight="主考" if i == 0 else "涉及",
                    )
                )
            saved += 1
        db.commit()
        return saved
    finally:
        db.close()


def _match_nodes(nodes: "Sequence[KnowledgeNode]", knowledge_point: str) -> list["KnowledgeNode"]:
    """考查知识点 → 图谱节点：精确标题优先，其次标题互含（近名容错）。"""
    if not knowledge_point:
        return []
    exact = [n for n in nodes if n.title == knowledge_point]
    if exact:
        return exact
    return [
        n for n in nodes if n.title and (n.title in knowledge_point or knowledge_point in n.title)
    ]


def render_exam(questions: list[dict], output_path: str | None = None) -> str:
    """渲染 .docx：试题 + 答案解析两部分，试题行标注考查知识点。"""
    if output_path is None:
        output_path = unique_output_path("exam", ".docx")
    doc = DocxDocument()
    style_document(doc)
    doc.add_heading("试卷", level=0)
    doc.add_heading("试题", level=1)
    for i, q in enumerate(questions, 1):
        knowledge_point = str(q.get("knowledge_point") or "").strip()
        tag = f"（考查知识点：{knowledge_point}）" if knowledge_point else ""
        doc.add_paragraph(f"{i}. [{q.get('type', '')}]{tag} {q.get('content', '')}")
        for option in q.get("options") or []:
            doc.add_paragraph(f"　　{option}")
    doc.add_page_break()
    doc.add_heading("答案解析", level=1)
    for i, q in enumerate(questions, 1):
        line = f"{i}. {q.get('answer', '')}"
        analysis = str(q.get("analysis") or "").strip()
        if analysis:
            line += f"　解析：{analysis}"
        doc.add_paragraph(line)
    doc.save(output_path)
    return output_path
