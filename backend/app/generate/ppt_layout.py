"""Wide-screen editorial layouts with bounded text and continuation pages."""

import math

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

WHITE = RGBColor.from_string("FFFFFF")
PAPER = RGBColor.from_string("F5F6F2")
MUTED = RGBColor.from_string("68746C")


def box(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def text(slide, value, x, y, w, h, size, color, bold=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    p = frame.paragraphs[0]
    p.space_after = Pt(0)
    p.line_spacing = 1.2
    run = p.add_run()
    run.text = value
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def pages(slides):
    """Break exceptionally long model output into readable pages without truncating it."""
    for data in slides:
        points = []
        for point in data.get("points", []):
            value = "；".join(map(str, point)) if isinstance(point, list) else str(point)
            points.extend(value[i : i + 72] for i in range(0, len(value), 72))
        title = str(data.get("title", ""))
        # Treat oversized titles as content, keeping all supplied text.
        if len(title) > 32:
            points = [title[i : i + 72] for i in range(0, len(title), 72)] + points
            title = title[:28] + "…"
        for start in range(0, max(1, len(points)), 4):
            yield {
                **data,
                "title": title + (" · 续" if start else ""),
                "role": data.get("role", "内容") if not start else "内容",
                "points": points[start : start + 4],
            }


def render_editorial(prs, slides, theme):
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    expanded = list(pages(slides))
    for index, data in enumerate(expanded):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = PAPER
        accent = theme["accent"]
        role = str(data.get("role", "内容")).removesuffix("页")
        title = str(data.get("title", ""))
        points = data.get("points", [])
        if role == "封面":
            box(slide, 9.65, 0, 3.69, 7.5, accent)
            text(slide, "EDUMIND  /  LESSON SERIES", 0.7, 0.65, 8, 0.4, 12, accent, True)
            text(
                slide,
                title,
                0.7,
                1.65,
                8.4,
                1.9,
                38 if len(title) > 15 else 46,
                theme["title"],
                True,
            )
            for n, point in enumerate(points):
                text(
                    slide,
                    point,
                    0.75,
                    3.75 + n * 0.55,
                    8.25,
                    0.55,
                    15 if len(point) > 36 else 18,
                    theme["text"],
                )
            text(slide, "01", 10.1, 2.1, 2.8, 1.8, 96, WHITE, True)
            text(slide, "从理解开始\n让知识发生连接", 10.1, 4.7, 2.6, 1.2, 19, WHITE)
        else:
            box(slide, 0.7, 0.6, 0.12, 0.28, accent)
            label = {"目录": "LEARNING PATH / 学习路径", "总结": "TAKEAWAYS / 课堂回顾"}.get(
                role, "EXPLORE / 知识探索"
            )
            text(slide, label, 0.98, 0.61, 10, 0.35, 11, MUTED, True)
            text(
                slide,
                title,
                0.7,
                1.18,
                11.9,
                1.1,
                30 if len(title) > 22 else 34,
                accent if role == "总结" else theme["title"],
                True,
            )
            if role == "目录":
                for n, point in enumerate(points):
                    y = 2.65 + n * 0.84
                    text(slide, f"{n + 1:02}", 0.8, y, 0.8, 0.55, 23, accent, True)
                    text(
                        slide,
                        point,
                        1.85,
                        y + 0.02,
                        10.5,
                        0.7,
                        19 if len(point) > 36 else 24,
                        theme["text"],
                    )
            else:
                count = len(points)
                cols = 2 if count > 1 else 1
                rows = max(1, math.ceil(count / cols))
                width = (11.93 - 0.3 * (cols - 1)) / cols
                height = (3.85 - 0.24 * (rows - 1)) / rows
                for n, point in enumerate(points):
                    x, y = 0.7 + n % cols * (width + 0.3), 2.55 + n // cols * (height + 0.24)
                    box(slide, x, y, width, height, WHITE)
                    text(slide, f"{n + 1:02}", x + 0.24, y + 0.16, 1, 0.35, 13, accent, True)
                    text(
                        slide,
                        point,
                        x + 0.24,
                        y + 0.65,
                        width - 0.48,
                        height - 0.75,
                        (17 if len(point) > 40 else 20) if rows > 1 else 25,
                        theme["text"],
                        role == "总结",
                    )
        text(slide, "EduMind  ·  教学课件", 0.72, 6.98, 8, 0.25, 9, MUTED)
        text(
            slide,
            f"{index + 1:02} / {len(expanded):02}",
            11.3,
            6.96,
            1.2,
            0.3,
            10,
            WHITE if role == "封面" else MUTED,
        )
