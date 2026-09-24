"""教学提纲生成：LLM 生成 Markdown 提纲 + python-docx 渲染。"""

import json
from functools import partial

from docx import Document as DocxDocument

from app.core.llm.base import ChatMessage
from app.core.llm.task_routing import get_llm_for
from app.generate import unique_output_path

# 「生成」任务：模型档位在设置页按任务选，未设置回落全局默认（CONTEXT.md「任务级模型」）。
# 入口仍叫 get_llm：既有测试用它替换对话能力（monkeypatch.setattr(本模块, "get_llm", ...)）。
get_llm = partial(get_llm_for, "generate")

_OUTLINE_PROMPT = """你是教学设计师。根据教学意图和知识内容，生成教学提纲（大纲）。

要求：分章节、分知识点，层次清晰，用 Markdown 格式（# 章节、## 小节、- 知识点）。

教学意图：__INTENT__
知识内容：__KNOWLEDGE__
"""


async def generate_outline(intent: dict, knowledge: str) -> str:
    llm = get_llm()
    prompt = _OUTLINE_PROMPT.replace("__INTENT__", json.dumps(intent, ensure_ascii=False)).replace(
        "__KNOWLEDGE__", knowledge[:6000]
    )
    result = await llm.chat([ChatMessage(role="user", content=prompt)])
    return result.content


def render_outline(text: str, output_path: str | None = None) -> str:
    """把 Markdown 提纲渲染为 .docx，返回文件路径（提纲作为生成物要有落盘文件可回看/下载）。

    `#` / `##` 映射为标题层级，`-` / `*` 映射为项目符号；其余行按段落处理。
    """
    if output_path is None:
        output_path = unique_output_path("outline", ".docx")

    doc = DocxDocument()
    doc.add_heading("教学提纲", level=0)
    for raw in (text or "").splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith(("- ", "* ")):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(line)

    doc.save(output_path)
    return output_path
