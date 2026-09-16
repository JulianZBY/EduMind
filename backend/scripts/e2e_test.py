"""EduMind 端到端验收：Playwright 驱动真实 Chrome，stub 网关全链路（不触外部服务）。

前置：后端 :8000（LLM_PROVIDER=stub / ASR_PROVIDER=stub）、前端 :5173（vite dev，/api 代理）。
运行：backend 目录下 `uv run python scripts/e2e_test.py`。
失败时截图到 e2e-artifacts/。
"""

import contextlib
import tempfile
import traceback
from pathlib import Path

from docx import Document
from playwright.sync_api import expect, sync_playwright

BASE = "http://localhost:5173"
ARTIFACTS = Path(__file__).resolve().parent.parent.parent / "e2e-artifacts"
RESULTS: list[tuple[str, bool, str]] = []


def scenario(name: str):
    """场景包装：记录结果，失败全页截图。"""

    def deco(fn):
        def run(page):
            try:
                fn(page)
                RESULTS.append((name, True, ""))
                print(f"  PASS  {name}")
            except Exception as exc:  # noqa: BLE001 — 测试跑包装器，任何场景异常都要记录并截图后继续
                ARTIFACTS.mkdir(parents=True, exist_ok=True)
                with contextlib.suppress(Exception):
                    page.screenshot(path=str(ARTIFACTS / f"FAIL_{name}.png"), full_page=True)
                RESULTS.append((name, False, traceback.format_exc(limit=5)))
                print(f"  FAIL  {name}: {exc}")
                print(traceback.format_exc(limit=5))

        return run

    return deco


@scenario("S1 冒烟：首页加载与四区结构")
def s1_smoke(page):
    page.goto(BASE, timeout=60000)
    expect(page.locator(".logo")).to_have_text("EduMind")
    expect(page.locator(".msg.assistant").first).to_contain_text("教学智能体", timeout=15000)
    tabs = page.locator(".tabs button")
    expect(tabs).to_have_count(4)
    for tab in ("知识库", "产物预览", "冲突审核", "会话信息"):
        expect(page.locator(".tabs button", has_text=tab)).to_be_visible()


@scenario("S2 知识库：上传讲义 → 解析入库 → 图谱成型")
def s2_upload(page):
    docx = Path(tempfile.gettempdir()) / "edumind_e2e_tcp讲义.docx"
    doc = Document()
    doc.add_heading("TCP 三次握手讲义", 0)
    doc.add_paragraph("TCP 三次握手：客户端发 SYN，服务端回 SYN+ACK，客户端再发 ACK。")
    doc.add_paragraph("释放连接需要四次挥手；滑动窗口机制实现流量控制。")
    doc.save(str(docx))

    page.locator("input[type=file]").set_input_files(str(docx))
    item = page.locator(".doc-item", has_text="edumind_e2e_tcp讲义.docx").first
    expect(item).to_be_visible(timeout=30000)
    expect(item.locator(".doc-status")).to_have_text("已完成", timeout=60000)
    expect(page.locator(".node-item").first).to_be_visible(timeout=30000)


@scenario("S3 备课对话：追问粒度 + 生成课件/教案/提纲")
def s3_chat(page):
    page.locator(".granularity button", has_text="精细").click()
    expect(page.locator(".granularity button.active")).to_have_text("精细")

    inp = page.locator("input[placeholder*='教学需求']")
    inp.fill("给大二学生讲 TCP 三次握手，45 分钟，学术风格")
    before = page.locator(".msg").count()
    page.locator("button", has_text="发送").click()

    # 备课完成 = 发送按钮恢复可用（sending 状态解除）
    expect(page.locator("button", has_text="备课中…")).to_have_count(0, timeout=120000)
    expect(page.locator("button", has_text="发送")).to_be_enabled()
    assert page.locator(".msg").count() >= before + 2, "对话轮数未增加"

    # 意图不完整时后端先追问：补全要素再答，直至产物下发（stub 最多两三轮）
    page.locator(".tabs button", has_text="产物预览").click()
    answer = "主题：TCP 三次握手；学段：大二；时长：45 分钟；风格：学术"
    for _ in range(3):
        if page.locator(".artifact-title", has_text="当前版本").count():
            break
        inp = page.locator("input[placeholder*='教学需求']")
        inp.fill(answer)
        page.locator("button", has_text="发送").click()
        expect(page.locator("button", has_text="备课中…")).to_have_count(0, timeout=120000)
        page.locator(".tabs button", has_text="产物预览").click()
    expect(page.locator(".artifact-title", has_text="当前版本")).to_be_visible()
    expect(page.locator("a.btn", has_text="PPT（")).to_be_visible()
    expect(page.locator("a.btn", has_text="Word 教案")).to_be_visible()


@scenario("S4 修改闭环（课件）：意见 → 重排 → 当前版本更新")
def s4_revise_ppt(page):
    rev = page.locator(".revise-area input")
    rev.fill("把要点再精简一些")
    page.locator(".revise-area button", has_text="修改").click()
    expect(rev).to_have_value("", timeout=60000)  # 成功后才清空输入


@scenario("S5 修改闭环（教案）：双入口第二入口")
def s5_revise_word(page):
    page.locator(".target-toggle button", has_text="教案").click()
    rev = page.locator(".revise-area input")
    expect(rev).to_have_attribute("placeholder", "修改教案：调整环节 / 增加活动 / 改作业")
    rev.fill("增加一个课堂提问环节")
    page.locator(".revise-area button", has_text="修改").click()
    expect(rev).to_have_value("", timeout=60000)


@scenario("S6 试卷：一键生成 → 题目预览 → docx 可下载")
def s6_exam(page):
    page.locator("button", has_text="生成试卷").click()
    expect(page.locator("summary", has_text="预览试卷（")).to_be_visible(timeout=120000)
    expect(page.locator(".exam-preview li").first).to_be_visible()
    expect(page.locator(".exam-preview .tag", has_text="考查知识点").first).to_contain_text(
        "考查知识点"
    )

    with page.expect_download(timeout=30000) as dl:
        page.locator("a.btn", has_text="试卷（").click()
    target = ARTIFACTS / dl.value.suggested_filename
    dl.value.save_as(str(target))
    assert target.stat().st_size > 1000, "试卷 docx 过小"


@scenario("S7 互动内容：一键生成 → 内联入口出现")
def s7_interactive(page):
    page.locator("button", has_text="生成互动内容").click()
    expect(page.locator("a.btn", has_text="互动内容").first).to_be_visible(timeout=120000)


@scenario("S8 会话持久化：刷新恢复消息/产物/粒度/会话名")
def s8_persist(page):
    page.reload(timeout=60000)
    expect(page.locator(".msg.assistant").first).to_contain_text("教学智能体")
    assert page.locator(".msg", has_text="给大二学生讲").count() >= 1, "教师消息未恢复"
    page.locator(".tabs button", has_text="产物预览").click()
    expect(page.locator(".artifact-title", has_text="当前版本")).to_be_visible()
    expect(page.locator(".granularity button.active")).to_have_text("精细")

    page.locator(".tabs button", has_text="会话信息").click()
    expect(page.locator(".panel strong").first).to_contain_text("握手")  # 首条消息自动命名


@scenario("S9 多会话：新建 / 切换 / 两段式删除")
def s9_sessions(page):
    page.locator(".session-wrap button.ghost", has_text="会话").click()
    page.locator(".session-new").click()
    expect(page.locator(".msg")).to_have_count(1, timeout=10000)  # 新会话仅欢迎语

    # 切回旧会话：菜单内点击旧会话条目
    page.locator(".session-wrap button.ghost", has_text="会话").click()
    items = page.locator(".session-item")
    expect(items).to_have_count(2)
    items.nth(1).locator(".session-main").click()
    expect(page.locator(".msg", has_text="给大二学生讲").first).to_be_visible(timeout=10000)

    # 删除当前（旧）会话：两段式确认（点击条目后菜单已关，先重新打开）
    page.locator(".session-wrap button.ghost", has_text="会话").click()
    expect(page.locator(".session-item.active")).to_have_count(1)
    page.locator(".session-item.active .session-delete").click()
    expect(page.locator(".session-item.active .session-delete")).to_have_text("确认删除？")
    page.locator(".session-item.active .session-delete").click()
    expect(page.locator(".msg")).to_have_count(1, timeout=10000)  # 回落到新会话
    expect(page.locator(".session-wrap button.ghost")).to_contain_text("会话（1）")


@scenario("S10 冲突审核：面板可用（全新库暂无冲突）")
def s10_conflicts(page):
    page.locator(".tabs button", has_text="冲突审核").click()
    expect(page.locator(".panel", has_text="冲突待审")).to_be_visible()


def main() -> int:
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    with sync_playwright() as p:
        browser = p.chromium.launch(channel="chrome", headless=True)
        page = browser.new_page(viewport={"width": 1440, "height": 900})
        page.set_default_timeout(15000)
        for fn in (
            s1_smoke,
            s2_upload,
            s3_chat,
            s4_revise_ppt,
            s5_revise_word,
            s6_exam,
            s7_interactive,
            s8_persist,
            s9_sessions,
            s10_conflicts,
        ):
            fn(page)
        browser.close()

    failed = [r for r in RESULTS if not r[1]]
    print(f"\n===== E2E 结果：{len(RESULTS) - len(failed)}/{len(RESULTS)} 通过 =====")
    for name, _ok, err in failed:
        print(f"  FAIL {name}\n{err}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
